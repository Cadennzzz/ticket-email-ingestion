"""
Sync the BL and SL sheets of Working tickets copy.xlsm into transactions.db
(source='excel' rows).

Each sheet row is keyed as before: BL by row number (excel-bl-<row>), SL by
its SID column (excel-sl-<SID>). For every sheet row:

  - no DB row with that key        -> ADD
  - same key, same event name      -> UPDATE any sheet fields that changed
  - same key, different event name -> FLAG, left alone (a row inserted or
    deleted in BL, or in SL's SEQUENCE-numbered SIDs 1-44, shifts keys onto
    a different sale; updating would overwrite one sale with another)

and every source='excel' DB row whose key is no longer in the sheet is
FLAGGED. Nothing is ever deleted. Rows are updated in place, so their ids,
and the matches / crosscheck links that point at them, are kept.

The old import only inserted missing keys, so sheet edits never reached
the DB (BL row 124 stayed $330 high).

Dry run by default: prints the plan and writes nothing. Back up
transactions.db, review the plan, then apply:

    python import_excel.py            # dry run
    python import_excel.py --apply

Read-only against the workbook: it is never saved or written back.
Column positions come from each sheet's header row (row 2), since the
layout has changed between versions.
"""

import argparse
import datetime
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from db import COLUMNS, get_connection

SOURCE_PATH = Path(__file__).parent / "Working tickets copy.xlsm"

HEADER_ROW = 2


def read_header_map(ws):
    """Map column name -> 0-based index from the sheet's header row."""
    rows_iter = ws.iter_rows(min_row=HEADER_ROW, values_only=True)
    header = next(rows_iter)
    col_map = {}
    for idx, name in enumerate(header):
        if name is not None:
            col_map[str(name).strip()] = idx
    return col_map, rows_iter


def get(row, col_map, name):
    idx = col_map.get(name)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def to_date_str(value):
    if value is None or value == "":
        return None
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    ts = pd.to_datetime(value, errors="coerce")
    if pd.isna(ts):
        return None
    return ts.strftime("%Y-%m-%d")


def to_float(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = str(value).replace("$", "").replace(",", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def to_int(value):
    f = to_float(value)
    return int(f) if f is not None else None


def to_transfer_status(value):
    if isinstance(value, bool):
        return "transferred" if value else None
    if value is None or value == "":
        return None
    return str(value).strip().lower()


# Fields taken from the sheet; anything else on an excel row is left alone.
SHEET_FIELDS = [
    "platform",
    "artist_or_event",
    "venue",
    "event_date",
    "purchase_date",
    "quantity",
    "price_per_ticket",
    "total_price",
    "transfer_status",
]


def _row(transaction_type, uid, **fields):
    data = {col: None for col in COLUMNS}
    data.update(
        is_ticket_transaction=True,
        transaction_type=transaction_type,
        needs_review=False,
        raw_email_uid=uid,
        source="excel",
        **fields,
    )
    return data


def bl_rows(ws):
    """(sheet row number, transaction dict) for every BL row with an Event."""
    col_map, rows_iter = read_header_map(ws)
    for row_num, row in enumerate(rows_iter, start=HEADER_ROW + 1):
        event = get(row, col_map, "Event")
        if event is None or str(event).strip() == "":
            continue
        yield row_num, _row(
            "buy",
            f"excel-bl-{row_num}",
            platform=get(row, col_map, "Account"),
            artist_or_event=str(event).strip(),
            venue=get(row, col_map, "Venue"),
            event_date=to_date_str(get(row, col_map, "Event Date")),
            purchase_date=to_date_str(get(row, col_map, "Date Purchased")),
            quantity=to_int(get(row, col_map, "Quantity")),
            price_per_ticket=to_float(get(row, col_map, "$ per Tix")),
            total_price=to_float(get(row, col_map, "Total Cost")),
        )


def sl_rows(ws):
    """(sheet row number, transaction dict) for every SL row with a SID."""
    col_map, rows_iter = read_header_map(ws)
    for row_num, row in enumerate(rows_iter, start=HEADER_ROW + 1):
        sid = get(row, col_map, "SID")
        if sid is None or str(sid).strip() == "":
            continue
        yield row_num, _row(
            "sell",
            f"excel-sl-{str(sid).strip()}",
            platform=get(row, col_map, "Marketplace"),
            artist_or_event=str(get(row, col_map, "Event") or "").strip() or None,
            venue=get(row, col_map, "Venue"),
            event_date=to_date_str(get(row, col_map, "Event Date")),
            purchase_date=to_date_str(get(row, col_map, "Date Sold")),
            quantity=to_int(get(row, col_map, "Tickets Sold")),
            price_per_ticket=to_float(get(row, col_map, "Sell Price (per ticket)")),
            total_price=to_float(get(row, col_map, "Gross Sale")),
            transfer_status=to_transfer_status(get(row, col_map, "Transferred?")),
        )


def _event_key(name) -> str:
    return " ".join(str(name or "").lower().split())


def _same(a, b) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) < 1e-6
    return a == b


def plan_sync(sheet_rows, db_rows):
    """
    Compare sheet rows [(row_num, data)] with existing source='excel' DB
    rows of one sheet ({uid: row dict}). Returns (adds, updates, flags):
    adds [(row_num, data)], updates [(row_num, db_row, {field: (old, new)})],
    flags [(row_num or None, uid, reason)].
    """
    adds, updates, flags = [], [], []
    seen = set()
    for row_num, data in sheet_rows:
        uid = data["raw_email_uid"]
        if uid in seen:
            flags.append((row_num, uid, "key appears twice in the sheet; second copy ignored"))
            continue
        seen.add(uid)
        current = db_rows.get(uid)
        if current is None:
            adds.append((row_num, data))
        elif _event_key(current["artist_or_event"]) != _event_key(data["artist_or_event"]):
            flags.append(
                (
                    row_num,
                    uid,
                    f"event changed: DB {current['artist_or_event']!r} "
                    f"(qty {current['quantity']}, ${current['total_price']}) vs sheet "
                    f"{data['artist_or_event']!r} (qty {data['quantity']}, ${data['total_price']})",
                )
            )
        else:
            changes = {
                f: (current[f], data[f]) for f in SHEET_FIELDS if not _same(current[f], data[f])
            }
            if changes:
                updates.append((row_num, current, changes))
    for uid in sorted(set(db_rows) - seen):
        r = db_rows[uid]
        flags.append(
            (
                None,
                uid,
                f"not in sheet any more: id {r['id']} {r['artist_or_event']!r} "
                f"{r['event_date']} qty {r['quantity']} ${r['total_price']}",
            )
        )
    return adds, updates, flags


def load_db_rows(conn, transaction_type):
    conn.row_factory = sqlite3.Row
    rows = conn.execute(
        "SELECT * FROM transactions WHERE source = 'excel' AND transaction_type = ?",
        (transaction_type,),
    ).fetchall()
    return {r["raw_email_uid"]: dict(r) for r in rows}


def print_plan(sheet, adds, updates, flags) -> None:
    print(f"\n=== {sheet}: {len(adds)} add, {len(updates)} update, {len(flags)} flag ===")
    for row_num, d in adds:
        print(
            f"  ADD    row {row_num:>4}  {d['raw_email_uid']:<14} {d['artist_or_event']!r} "
            f"{d['event_date']} qty {d['quantity']} ${d['total_price']}"
        )
    for row_num, cur, changes in updates:
        diff = "; ".join(f"{f}: {old!r} -> {new!r}" for f, (old, new) in changes.items())
        print(f"  UPDATE row {row_num:>4}  {cur['raw_email_uid']:<14} id {cur['id']} {cur['artist_or_event']!r}: {diff}")
    for row_num, uid, reason in flags:
        where = f"row {row_num:>4}" if row_num else "no row  "
        print(f"  FLAG   {where}  {uid:<14} {reason}")


def apply_plan(conn, adds, updates) -> None:
    column_list = ", ".join(COLUMNS)
    placeholders = ", ".join("?" for _ in COLUMNS)
    for _, data in adds:
        conn.execute(
            f"INSERT INTO transactions ({column_list}) VALUES ({placeholders})",
            [data[c] for c in COLUMNS],
        )
    for _, cur, changes in updates:
        sets = ", ".join(f"{f} = ?" for f in changes)
        conn.execute(
            f"UPDATE transactions SET {sets} WHERE id = ?",
            [new for _, new in changes.values()] + [cur["id"]],
        )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--apply", action="store_true", help="write the planned adds/updates")
    args = parser.parse_args()

    wb = load_workbook(SOURCE_PATH, keep_vba=True, read_only=True, data_only=True)
    try:
        sheets = {"BL": ("buy", list(bl_rows(wb["BL"]))), "SL": ("sell", list(sl_rows(wb["SL"])))}
    finally:
        wb.close()

    conn = get_connection()
    try:
        plans = {}
        for sheet, (ttype, rows) in sheets.items():
            plans[sheet] = plan_sync(rows, load_db_rows(conn, ttype))
            print_plan(sheet, *plans[sheet])

        if not args.apply:
            print("\nDry run: nothing written. Re-run with --apply to write adds and updates.")
            return
        with conn:
            for adds, updates, _ in plans.values():
                apply_plan(conn, adds, updates)
        n_add = sum(len(a) for a, _, _ in plans.values())
        n_upd = sum(len(u) for _, u, _ in plans.values())
        print(f"\nApplied: {n_add} added, {n_upd} updated. Flagged rows were not touched.")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
