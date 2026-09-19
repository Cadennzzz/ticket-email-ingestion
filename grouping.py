"""
Shared grouping/normalization helpers for pending scraped transactions.

Used by pending_groups.py and dashboard.py to turn raw source='email',
promoted=0 rows into deduplicated, aggregated groups suitable for manual
copy-paste into Working tickets copy.xlsm's BL/SL sheets. normalize_tier
also lives here (rather than inline in one script) so import_excel.py can
reuse it later if ticket-type normalization is ever needed on the
historical-import side.

This is a best-effort heuristic, not a guarantee — always eyeball the
suggested groups before pasting anything into the workbook.
"""

import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from match import normalize_event

SOURCE_PATH = Path(__file__).parent / "Working tickets copy.xlsm"
HEADER_ROW = 2

# "a few percent apart" — same tier, price differs by more than this
# fraction of the lower price and it gets split into its own group.
PRICE_DIVERGENCE_THRESHOLD = 0.05
PRICE_DIVERGENCE_FLAG = "same tier, different price — verify"

UNSPECIFIED_TIER = "UNSPECIFIED"

_ROMAN_NUMERALS = {"I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"}
_QUALIFIER_WORDS = {"PLATINUM", "GOLD", "SILVER", "PACKAGE", "UPGRADE", "EXPERIENCE", "ADDON"}
_TIER_ALIASES = {
    "GENERAL ADMISSION": "GA",
    "GENERAL ADMISION": "GA",  # common misspelling
    "GA": "GA",
    "VIP": "VIP",
}


def normalize_tier(raw) -> str:
    """
    Normalize a raw ticket_type string into a coarse tier for grouping:
    strip trailing digits / roman numerals / qualifier words (PLATINUM,
    GOLD, ...), then map known variants (GENERAL ADMISSION, GA, GA1,
    GA 2, ... -> GA; VIP, VIP1, VIP PLATINUM, ... -> VIP). Anything else
    is kept as its own normalized (trimmed/uppercased) literal.
    """
    if raw is None:
        return UNSPECIFIED_TIER
    s = re.sub(r"[^A-Za-z0-9\s\-]", " ", str(raw)).strip().upper()
    if not s:
        return UNSPECIFIED_TIER

    tokens = s.split()
    changed = True
    while changed and tokens:
        changed = False
        last = tokens[-1]

        if last.isdigit() and len(tokens) > 1:
            tokens.pop()
            changed = True
            continue
        if last in _ROMAN_NUMERALS and len(tokens) > 1:
            tokens.pop()
            changed = True
            continue
        if last in _QUALIFIER_WORDS and len(tokens) > 1:
            tokens.pop()
            changed = True
            continue

        m = re.match(r"^([A-Z]+)(\d+)$", last)
        if m:
            tokens[-1] = m.group(1)
            changed = True
            continue

    normalized = " ".join(t for t in tokens if t).strip()
    if not normalized:
        normalized = s

    return _TIER_ALIASES.get(normalized, normalized)


def _split_by_price(rows, threshold=PRICE_DIVERGENCE_THRESHOLD):
    """
    Cluster rows by price_per_ticket: start a new cluster once a row's
    price diverges from its cluster's reference (first) price by more
    than `threshold`. Rows with no price at all can't be compared, so
    they're attached to the first cluster.
    """
    priced = [r for r in rows if r.get("price_per_ticket") is not None]
    unpriced = [r for r in rows if r.get("price_per_ticket") is None]

    if not priced:
        return [rows] if rows else []

    priced_sorted = sorted(priced, key=lambda r: r["price_per_ticket"])
    clusters = [[priced_sorted[0]]]
    for r in priced_sorted[1:]:
        ref = clusters[-1][0]["price_per_ticket"]
        price = r["price_per_ticket"]
        diverges = bool(ref) and abs(price - ref) / ref > threshold
        if diverges:
            clusters.append([r])
        else:
            clusters[-1].append(r)

    if unpriced:
        clusters[0].extend(unpriced)

    return clusters


def _seat_fields_compatible(a, b) -> bool:
    for field in ("section", "row", "seat"):
        av, bv = a.get(field), b.get(field)
        if av is not None and bv is not None and av != bv:
            return False
    return True


def _is_seat_wildcard(r) -> bool:
    return r.get("section") is None and r.get("row") is None and r.get("seat") is None


def _split_by_seat_fields(rows):
    """
    Split rows into sub-groups where section/row/seat don't conflict.
    None is a wildcard, not a distinct value — only actual differing
    non-null values force a split.

    Rows with ALL THREE fields null are excluded from the compatibility
    clustering itself and attached afterward to the largest resulting
    subgroup. Otherwise a naive pairwise union-find lets a single fully-
    null row transitively "bridge" two rows that actually conflict (e.g.
    section=101 and section=102 both count as compatible with a
    section=None row, so plain union-find would merge 101 and 102
    together through it) — excluding wildcards from the clustering step
    avoids that. Rows with a *partial* null pattern (e.g. section set,
    row null) can still bridge in rarer cases; treated as an acceptable
    heuristic limitation given this is a presentation aid, not a solver.
    """
    n = len(rows)
    if n <= 1:
        return [rows] if rows else []

    concrete = [r for r in rows if not _is_seat_wildcard(r)]
    wildcard = [r for r in rows if _is_seat_wildcard(r)]

    if not concrete:
        return [rows]

    m = len(concrete)
    parent = list(range(m))

    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(x, y):
        rx, ry = find(x), find(y)
        if rx != ry:
            parent[rx] = ry

    for i in range(m):
        for j in range(i + 1, m):
            if _seat_fields_compatible(concrete[i], concrete[j]):
                union(i, j)

    buckets = defaultdict(list)
    for i in range(m):
        buckets[find(i)].append(concrete[i])
    groups = list(buckets.values())

    if wildcard:
        groups.sort(key=len, reverse=True)
        groups[0].extend(wildcard)

    return groups


def _summarize(rows, flag=None):
    total_quantity = sum(r.get("quantity") or 0 for r in rows)
    total_cost = sum(r["total_price"] for r in rows if r.get("total_price") is not None)
    avg_price = (total_cost / total_quantity) if total_quantity else None

    representative = rows[0]
    return {
        "artist_or_event": representative.get("artist_or_event"),
        "venue": representative.get("venue"),
        "event_date": representative.get("event_date"),
        "tier": representative.get("_tier"),
        "total_quantity": total_quantity,
        "total_cost": total_cost,
        "avg_price_per_ticket": avg_price,
        "contributing_ids": [r.get("id") for r in rows],
        "contributing_uids": [r.get("raw_email_uid") for r in rows],
        "rows": rows,
        "flag": flag,
    }


def group_pending_rows(rows):
    """
    Group raw pending (source='email', promoted=0) transaction dicts into
    aggregated blocks suitable for a single Excel row each.

    Each input row dict must use real None (not NaN) for missing values,
    and have at least: id, artist_or_event, venue, event_date,
    ticket_type, price_per_ticket, quantity, total_price, section, row,
    seat, raw_email_uid.

    Grouping key is (normalized event name, event_date, normalized tier).
    Within a bucket, rows are further split (never silently merged) when
    price diverges meaningfully for the same tier, or when section/row/
    seat conflict. Returned groups are sorted by event_date then event
    name for stable, chronological output.
    """
    buckets = defaultdict(list)
    for r in rows:
        r = dict(r)
        r["_tier"] = normalize_tier(r.get("ticket_type"))
        key = (normalize_event(r.get("artist_or_event")), r.get("event_date"), r["_tier"])
        buckets[key].append(r)

    groups = []
    for bucket_rows in buckets.values():
        price_clusters = _split_by_price(bucket_rows)
        flag = PRICE_DIVERGENCE_FLAG if len(price_clusters) > 1 else None

        for cluster in price_clusters:
            for seat_group in _split_by_seat_fields(cluster):
                groups.append(_summarize(seat_group, flag=flag))

    groups.sort(key=lambda g: (g["event_date"] or "", g["artist_or_event"] or "", g["tier"] or ""))
    return groups


def read_existing_event_names(source_path=SOURCE_PATH):
    """Read-only: collect every non-empty 'Event' value from BL and SL in
    Working tickets copy.xlsm, for collision-checking suggested names."""
    names = set()
    wb = load_workbook(source_path, keep_vba=True, read_only=True, data_only=True)
    try:
        for sheet_name in ("BL", "SL"):
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(min_row=HEADER_ROW, values_only=True)
            header = next(rows_iter)
            col_map = {}
            for idx, name in enumerate(header):
                if name is not None:
                    col_map[str(name).strip()] = idx
            event_idx = col_map.get("Event")
            if event_idx is None:
                continue
            for row in rows_iter:
                if event_idx < len(row) and row[event_idx]:
                    names.add(str(row[event_idx]).strip())
    finally:
        wb.close()
    return names


def suggest_event_name(base_name, event_date, tier, existing_names) -> str:
    """Suggest a unique Event name, disambiguating with date then tier,
    then a numeric suffix as a last resort."""
    base_name = (base_name or "Unknown Event").strip()
    if base_name not in existing_names:
        return base_name

    candidate = f"{base_name} {event_date}" if event_date else base_name
    if candidate not in existing_names:
        return candidate

    candidate2 = f"{candidate} {tier}" if tier and tier != UNSPECIFIED_TIER else candidate
    if candidate2 not in existing_names:
        return candidate2

    i = 2
    while True:
        numbered = f"{candidate2} ({i})"
        if numbered not in existing_names:
            return numbered
        i += 1
