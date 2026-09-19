"""
One-off diagnostic: compare Net Profit from Working tickets copy.xlsm's SL
sheet against match.py's computed net_profit in the matches table.

Read-only against both the workbook and transactions.db. No writes.

Run with:
    python diagnose_profit_gap.py
"""

from pathlib import Path

from openpyxl import load_workbook

from db import get_connection

SOURCE_PATH = Path(__file__).parent / "Working tickets copy.xlsm"
HEADER_ROW = 2


def read_header_map(ws):
    rows_iter = ws.iter_rows(min_row=HEADER_ROW, values_only=True)
    header = next(rows_iter)
    col_map = {}
    for idx, name in enumerate(header):
        if name is not None:
            col_map[str(name).strip()] = idx
    return col_map, rows_iter


def get(row, col_map, name):
    idx = col_map.get(name)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def to_float(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace("$", "").replace(",", "").strip())
    except ValueError:
        return None


def read_sl_rows():
    wb = load_workbook(SOURCE_PATH, keep_vba=True, read_only=True, data_only=True)
    try:
        ws = wb["SL"]
        col_map, rows_iter = read_header_map(ws)
        rows = []
        for row in rows_iter:
            sid = get(row, col_map, "SID")
            if sid is None or str(sid).strip() == "":
                continue
            rows.append(
                {
                    "sid": sid,
                    "event": get(row, col_map, "Event"),
                    "venue": get(row, col_map, "Venue"),
                    "event_date": get(row, col_map, "Event Date"),
                    "net_profit": to_float(get(row, col_map, "Net Profit")),
                    "gross_sale": to_float(get(row, col_map, "Gross Sale")),
                    "purchase_price_per_ticket": to_float(
                        get(row, col_map, "Purchase Price (per ticket)")
                    ),
                    "tickets_sold": to_float(get(row, col_map, "Tickets Sold")),
                }
            )
        return rows
    finally:
        wb.close()


def main() -> None:
    sl_rows = read_sl_rows()

    sl_net_profit_total = sum(r["net_profit"] for r in sl_rows if r["net_profit"] is not None)
    sl_gross_sale_total = sum(r["gross_sale"] for r in sl_rows if r["gross_sale"] is not None)
    sl_implied_cost_total = sum(
        r["purchase_price_per_ticket"] * r["tickets_sold"]
        for r in sl_rows
        if r["purchase_price_per_ticket"] is not None and r["tickets_sold"] is not None
    )

    conn = get_connection()
    try:
        matches_net_profit_total = conn.execute(
            "SELECT COALESCE(SUM(net_profit), 0) FROM matches"
        ).fetchone()[0]

        uid_to_id = dict(
            conn.execute(
                "SELECT raw_email_uid, id FROM transactions WHERE raw_email_uid LIKE 'excel-sl-%'"
            ).fetchall()
        )

        matched_sell_ids = {
            row[0] for row in conn.execute("SELECT DISTINCT sell_id FROM matches").fetchall()
        }
    finally:
        conn.close()

    print("--- Net Profit comparison ---")
    print(f"Working tickets copy.xlsm SL 'Net Profit' total: ${sl_net_profit_total:,.2f}")
    print(f"matches table net_profit total:                  ${matches_net_profit_total:,.2f}")
    print(f"Difference (xlsm - matches):                     ${sl_net_profit_total - matches_net_profit_total:,.2f}")
    print()
    print(f"SL 'Gross Sale' total:                            ${sl_gross_sale_total:,.2f}")
    print(f"SL implied cost (Purchase Price/tix * Tickets Sold): ${sl_implied_cost_total:,.2f}")
    print(f"Gross Sale - implied cost:                        ${sl_gross_sale_total - sl_implied_cost_total:,.2f}")
    print("(compare that line to the xlsm Net Profit total above — if they don't match,")
    print(" the sheet's Net Profit already nets out something match.py doesn't, e.g. fees.)")

    print()
    print("--- SIDs with Net Profit populated but not reflected in matches ---")
    flagged = 0
    for r in sl_rows:
        if r["net_profit"] is None:
            continue
        uid = f"excel-sl-{str(r['sid']).strip()}"
        tid = uid_to_id.get(uid)
        if tid is None:
            print(
                f"  SID={r['sid']} event={r['event']!r} venue={r['venue']!r} "
                f"date={r['event_date']} -> sell row not found in transactions.db (uid={uid})"
            )
            flagged += 1
        elif tid not in matched_sell_ids:
            print(
                f"  SID={r['sid']} event={r['event']!r} venue={r['venue']!r} "
                f"date={r['event_date']} -> transactions.id={tid} has no match in the matches table"
            )
            flagged += 1
    if not flagged:
        print("  (none)")


if __name__ == "__main__":
    main()
